# -*- coding: utf-8 -*-
"""
个人网页后端服务
- 读取同目录下的 information.xlsx（每次请求都重新读取，做到自适应表格修改）
- 将每个 sheet 解析为：每列 = 一个模块；第一行 = 模块名；其余行 = 图文内容
- 若某行文本与文件夹中某图片名（不含扩展名）一致，则渲染为图片，否则渲染为文字
- 自动将"出国照片"文件夹中的图片作为单独切页呈现
- 提供首页、/api/data 数据接口、/images/<文件> 图片服务
"""
import os
from flask import Flask, send_from_directory, jsonify, abort

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "information.xlsx")
ABROAD_DIR = os.path.join(BASE, "出国照片")  # 出国照片文件夹

app = Flask(__name__, static_folder=BASE, static_url_path="")

IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg")


def list_images():
    """返回 BASE 文件夹中所有图片文件名列表（不含子文件夹）"""
    try:
        files = [f for f in os.listdir(BASE) if f.lower().endswith(IMAGE_EXTS)]
    except FileNotFoundError:
        files = []
    return sorted(files)


def image_base_map():
    """{图片名(无扩展名): 实际文件名}。重名时保留第一个。"""
    m = {}
    for f in list_images():
        base = os.path.splitext(f)[0]
        m.setdefault(base, f)
    return m


def get_abroad_photos():
    """读取"出国照片"文件夹中的所有图片，作为单独切页的内容"""
    photos = []
    if not os.path.exists(ABROAD_DIR):
        return photos
    for f in sorted(os.listdir(ABROAD_DIR)):
        if not f.lower().endswith(IMAGE_EXTS):
            continue
        base = os.path.splitext(f)[0]
        photos.append({
            "type": "image",
            "label": base,
            "src": "/images/出国照片/" + f
        })
    return photos


@app.route("/")
def index():
    """返回首页 index.html"""
    idx = os.path.join(BASE, "index.html")
    if not os.path.exists(idx):
        return "index.html 不存在", 404
    return send_from_directory(BASE, "index.html")


@app.route("/api/data")
def api_data():
    """
    返回结构：
    {
      "sheets": [
        { "name": "...", "modules": [ {"title": "...", "items": [ {"type":"text"|"image", ...} ]} ] }
      ],
      "images": [...],
      "abroad": [ {"label": "...", "src": "..."} ]
    }
    每次请求都重新读取 xlsx，从而内容随表格修改而变化。
    """
    abroad_photos = get_abroad_photos()
    sheets = []

    if not os.path.exists(XLSX):
        # 即便 xlsx 不存在，也保留"出国照片"切页
        if abroad_photos:
            sheets.append({
                "name": "出国照片",
                "modules": [{"title": "海外行迹", "items": abroad_photos}]
            })
        return jsonify({
            "error": "information.xlsx 不存在",
            "sheets": sheets,
            "images": list_images(),
            "abroad": abroad_photos
        }), 200

    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    img_map = image_base_map()

    for ws in wb.worksheets:
        modules = []
        max_col = ws.max_column or 0
        for col_idx in range(1, max_col + 1):
            cells = [ws.cell(row=r, column=col_idx).value for r in range(1, (ws.max_row or 1) + 1)]
            # 去除末尾空值
            while cells and (cells[-1] is None or (isinstance(cells[-1], str) and not cells[-1].strip())):
                cells.pop()
            if not cells:
                continue
            title = str(cells[0]).strip() if cells[0] is not None else ""
            items = []
            for c in cells[1:]:
                if c is None:
                    continue
                text = str(c).strip()
                if not text:
                    continue
                img_file = img_map.get(text)
                if img_file:
                    items.append({"type": "image", "label": text, "src": "/images/" + img_file})
                else:
                    items.append({"type": "text", "value": text})
            modules.append({"title": title, "items": items})
        sheets.append({"name": ws.title, "modules": modules})

    # 将"出国照片"文件夹作为单独切页追加到末尾
    if abroad_photos:
        sheets.append({
            "name": "出国照片",
            "modules": [{"title": "海外行迹", "items": abroad_photos}]
        })

    return jsonify({
        "sheets": sheets,
        "images": list_images(),
        "abroad": abroad_photos
    })


@app.route("/images/<path:filename>")
def images(filename):
    """支持子目录路径，如 /images/出国照片/冰岛.jpg"""
    return send_from_directory(BASE, filename)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8000, debug=True)
